import io
import json

import docx

from app.knowledge.service import Section, chunk_sections
from app.llm.base import ToolCall
from app.llm.fake import Scripted
from app.tools import registry
from tests.helpers import chat, class_with_student, reply_text

LESSON = (
    "Chương 3. Phương trình bậc nhất một ẩn. Phương trình bậc nhất một ẩn có dạng ax + b = 0 với a khác 0. "
    "Để giải, ta chuyển vế và chia hai vế cho a.\n\n"
    "Quang hợp là quá trình cây xanh dùng ánh sáng mặt trời để tạo chất hữu cơ từ khí cacbonic và nước."
)


async def upload(client, class_id: str, name: str, data: bytes, content_type: str, title: str = ""):
    return await client.post(
        f"/api/classes/{class_id}/documents",
        files={"file": (name, data, content_type)},
        data={"title": title},
    )


async def test_teacher_uploads_text_document_and_it_is_chunked(as_user) -> None:
    teacher, student, _, room = await class_with_student(as_user)
    response = await upload(teacher, room["id"], "chuong3.txt", LESSON.encode(), "text/plain", "Bài giảng chương 3")
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "processing"
    # Background tasks run before the test client returns.
    documents = (await teacher.get(f"/api/classes/{room['id']}/documents")).json()
    assert documents[0]["status"] == "ready" and documents[0]["chunk_count"] >= 1
    chunks = (await teacher.get(f"/api/classes/{room['id']}/documents/{documents[0]['id']}/chunks")).json()
    assert "Phương trình bậc nhất" in chunks[0]["content"]
    # Students may list documents but not upload.
    assert len((await student.get(f"/api/classes/{room['id']}/documents")).json()) == 1
    denied = await upload(student, room["id"], "x.txt", b"hello world text", "text/plain")
    assert denied.status_code == 403


async def test_docx_upload_keeps_headings_as_locations(as_user) -> None:
    teacher, _, _, room = await class_with_student(as_user)
    document = docx.Document()
    document.add_heading("Bài 1: Phân số", level=1)
    document.add_paragraph("Phân số tối giản là phân số mà tử và mẫu không còn ước chung nào khác 1.")
    buffer = io.BytesIO()
    document.save(buffer)
    response = await upload(
        teacher,
        room["id"],
        "phan-so.docx",
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 201
    doc_id = response.json()["id"]
    chunks = (await teacher.get(f"/api/classes/{room['id']}/documents/{doc_id}/chunks")).json()
    assert chunks[0]["location"] == "Bài 1: Phân số"


async def test_unsupported_and_empty_files_are_rejected(as_user) -> None:
    teacher, _, _, room = await class_with_student(as_user)
    assert (await upload(teacher, room["id"], "virus.exe", b"MZ...", "application/octet-stream")).status_code == 415
    assert (await upload(teacher, room["id"], "empty.txt", b"", "text/plain")).status_code == 400


async def test_class_doc_search_returns_relevant_chunk_only_from_own_class(as_user, fake_llm) -> None:
    teacher, student, _, room = await class_with_student(as_user)
    await upload(teacher, room["id"], "bai.txt", LESSON.encode(), "text/plain", "Bài giảng")
    other_teacher = await as_user("teacher")
    other = (await other_teacher.post("/api/classes", json={"name": "Lớp khác", "subject": "sinh", "grade": 7})).json()
    await upload(
        other_teacher, other["id"], "secret.txt", b"Bi mat lop khac: quang hop quang hop quang hop", "text/plain"
    )

    fake_llm.script(
        Scripted(tool_calls=[ToolCall(id="s1", name="class_doc_search", arguments={"query": "quang hợp là gì"})]),
        "Theo tài liệu lớp, quang hợp là ...",
    )
    events = await chat(student, "Quang hợp là gì?", class_id=room["id"])
    assert reply_text(events).endswith("Theo tài liệu lớp, quang hợp là ...")
    result = json.loads(fake_llm.calls[-1]["messages"][-1].content)
    titles = {hit["title"] for hit in result["results"]}
    assert titles == {"Bài giảng"}
    assert "Quang hợp" in result["results"][0]["content"]


async def test_deleting_document_removes_it(as_user) -> None:
    teacher, _, _, room = await class_with_student(as_user)
    doc = (await upload(teacher, room["id"], "a.txt", LESSON.encode(), "text/plain")).json()
    assert (await teacher.delete(f"/api/classes/{room['id']}/documents/{doc['id']}")).status_code == 204
    assert (await teacher.get(f"/api/classes/{room['id']}/documents")).json() == []


def test_chunking_respects_size_and_overlap() -> None:
    text = " ".join(f"Câu số {index} nói về phương trình." for index in range(200))
    chunks = chunk_sections([Section(text, "trang 1")], size=300, overlap=50)
    assert len(chunks) > 5
    assert all(len(chunk.text) <= 360 for chunk in chunks)
    assert all(chunk.location == "trang 1" for chunk in chunks)


def _ctx(role: str = "student", tools=("math_check",), class_id: str | None = "c1") -> registry.ToolContext:
    return registry.ToolContext(
        user_id="u1",
        role=role,
        class_id=class_id,
        grade=8,
        conversation_id=None,
        agent_slug="t",
        allowed_tools=frozenset(tools),
    )


async def test_math_check_operations() -> None:
    async def run(**arguments):
        outcome = await registry.execute(ToolCall(id="m", name="math_check", arguments=arguments), _ctx())
        return outcome, json.loads(outcome.content)

    _, solved = await run(operation="solve", expression="x^2 - 5x + 6 = 0")
    assert sorted(solved["solutions"]) == ["2", "3"]
    _, wrong = await run(operation="check_answer", expression="2x + 3 = 11", answer="x = 5")
    assert wrong["is_correct"] is False
    _, simplified = await run(operation="simplify", expression="(x^2 - 1)/(x - 1)")
    assert simplified["simplified"] == "x + 1"
    _, same = await run(operation="equivalent", expression="(x+1)^2", answer="x^2 + 2x + 1")
    assert same["equivalent"] is True


async def test_math_check_rejects_code_injection() -> None:
    for payload in ["__import__('os').system('dir')", "x.__class__", "open('a')", "lambda: 1"]:
        outcome = await registry.execute(
            ToolCall(id="m", name="math_check", arguments={"operation": "simplify", "expression": payload}), _ctx()
        )
        assert outcome.is_error, payload


async def test_tool_guard_checks_role_class_and_arguments() -> None:
    parent_quiz = await registry.execute(
        ToolCall(id="q", name="quiz_generator", arguments={"topic": "phân số"}),
        _ctx("parent", tools=("quiz_generator",)),
    )
    assert parent_quiz.denied == "Vai trò của người dùng không được dùng công cụ này."
    no_class = await registry.execute(
        ToolCall(id="d", name="class_doc_search", arguments={"query": "abc"}),
        _ctx(tools=("class_doc_search",), class_id=None),
    )
    assert no_class.denied == "Công cụ này chỉ dùng trong một lớp học."
    bad_args = await registry.execute(ToolCall(id="m", name="math_check", arguments={"operation": "fly"}), _ctx())
    assert bad_args.is_error and bad_args.denied is None and "Tham số không hợp lệ" in bad_args.content
    unknown = await registry.execute(ToolCall(id="x", name="shell", arguments={}), _ctx(tools=("shell",)))
    assert unknown.denied == "Công cụ không có trong danh mục của nền tảng."


async def test_quiz_hides_answers_from_students(fake_llm) -> None:
    quiz = '{"questions":[{"type":"mcq","question":"2+2?","options":["A. 3","B. 4"],"answer":"B","explanation":"vì"}]}'
    fake_llm.script(quiz, quiz)
    student = await registry.execute(
        ToolCall(id="q", name="quiz_generator", arguments={"topic": "cộng", "with_answers": True}),
        _ctx("student", tools=("quiz_generator",), class_id=None),
    )
    teacher = await registry.execute(
        ToolCall(id="q", name="quiz_generator", arguments={"topic": "cộng", "with_answers": True}),
        _ctx("teacher", tools=("quiz_generator",), class_id=None),
    )
    assert "answer" not in json.loads(student.content)["questions"][0]
    assert json.loads(teacher.content)["questions"][0]["answer"] == "B"
    assert student.usage and student.usage[0].provider == "fake"


async def test_rubric_grader_gives_students_feedback_without_scores(fake_llm) -> None:
    graded = json.dumps(
        {
            "criteria": [
                {"name": "Bố cục", "score": 1.5, "max_score": 2, "evidence": "...", "suggestion": "Thêm kết bài"}
            ],
            "total": 1.5,
            "max_total": 2,
            "strengths": ["Mở bài hay"],
            "improvements": ["Thiếu kết bài"],
            "injection_detected": False,
        }
    )
    fake_llm.script(graded)
    outcome = await registry.execute(
        ToolCall(id="r", name="rubric_grader", arguments={"essay": "Mùa thu đến rồi, lá vàng rơi đầy sân trường em."}),
        _ctx("student", tools=("rubric_grader",)),
    )
    result = json.loads(outcome.content)
    assert "total" not in result and "score" not in result["criteria"][0]
    assert result["criteria"][0]["suggestion"] == "Thêm kết bài"


async def test_tools_endpoint_lists_registry(as_user) -> None:
    teacher = await as_user("teacher")
    tools = (await teacher.get("/api/tools")).json()
    assert {tool["name"] for tool in tools} == set(registry.TOOLS)
    math = next(tool for tool in tools if tool["name"] == "math_check")
    assert "ctx" not in math["parameters"]["properties"]
